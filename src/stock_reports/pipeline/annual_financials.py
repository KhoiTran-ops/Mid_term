"""Import the reviewed CafeF annual workbook without mixing it into quarterly facts."""

from datetime import datetime, UTC
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3

from openpyxl import load_workbook


KEY_HEADERS = {'firm': 'symbol', 'year': 'year'}
COMPANY_HEADERS = {'establish', 'total employee', 'restate', 'note'}
MARKET_HEADERS = {'số lượng cổ phiếu lưu hành', 'giá đóng cửa'}
MISSING_TEXT = {'', '-', '--', 'n/a', 'na', 'none', 'null'}


def clean_header(value: object) -> str:
    text = ' '.join(str(value or '').split())
    # Accounting form line codes vary by sheet but do not change metric meaning.
    return re.sub(r'\s*\((?:10|50|60)\)$', '', text).strip()


def raw_and_numeric(value: object) -> tuple[str, float | None]:
    if value is None:
        return '', None
    if isinstance(value, bool):
        return str(value), None
    if isinstance(value, (int, float)):
        number = float(value)
        return str(value), number if math.isfinite(number) else None
    raw = str(value).strip()
    return raw, None


def metric_category(metric: str) -> str:
    folded = metric.casefold()
    if folded in COMPANY_HEADERS:
        return 'company'
    if folded in MARKET_HEADERS:
        return 'market'
    return 'financial'


def import_annual_workbook(database: Path, workbook: Path, *, source_url: str) -> dict:
    workbook = Path(workbook).resolve()
    if not workbook.is_file() or workbook.suffix.lower() != '.xlsx':
        raise ValueError('Annual source must be an existing .xlsx workbook')
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    imported_at = datetime.now(UTC).isoformat()
    connection = sqlite3.connect(database)
    connection.execute('PRAGMA journal_mode=WAL')
    try:
        existing = connection.execute(
            'SELECT quality_json FROM annual_imports WHERE workbook_sha256=?', (digest,)
        ).fetchone()
        if existing:
            quality = json.loads(existing[0])
            quality.update(status='already_imported', workbook_sha256=digest)
            return quality

        book = load_workbook(workbook, read_only=True, data_only=True)
        workbook_rows = saved = duplicates = conflicts = 0
        symbols: set[str] = set()
        years: set[int] = set()
        schemas: dict[str, list[str]] = {}
        with connection:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = [clean_header(value) for value in next(rows, ())]
                schemas[sheet.title] = headers
                folded = [header.casefold() for header in headers]
                try:
                    symbol_index = folded.index('firm')
                    year_index = folded.index('year')
                except ValueError as error:
                    raise ValueError(f'Missing Firm/Year columns in sheet {sheet.title}') from error
                for source_row, values in enumerate(rows, 2):
                    workbook_rows += 1
                    symbol = str(values[symbol_index] or '').strip().upper()
                    try:
                        year = int(values[year_index])
                    except (TypeError, ValueError):
                        continue
                    if not symbol or not 1900 <= year <= 2100:
                        continue
                    symbols.add(symbol)
                    years.add(year)
                    for index, header in enumerate(headers):
                        if not header or header.casefold() in KEY_HEADERS or header.lower().startswith('unnamed:'):
                            continue
                        value = values[index] if index < len(values) else None
                        raw, numeric = raw_and_numeric(value)
                        if raw.casefold() in MISSING_TEXT:
                            continue
                        current = connection.execute(
                            'SELECT value_raw,value_numeric FROM annual_financial_facts '
                            'WHERE symbol=? AND fiscal_year=? AND metric=?',
                            (symbol, year, header),
                        ).fetchone()
                        if current:
                            same_numeric = current[1] == numeric or (
                                current[1] is not None and numeric is not None and
                                math.isclose(float(current[1]), numeric, rel_tol=1e-12)
                            )
                            if current[0] == raw or same_numeric:
                                duplicates += 1
                            else:
                                conflicts += 1
                                connection.execute(
                                    'INSERT INTO data_anomalies(dataset,record_key,reason,raw_json,detected_at) '
                                    'VALUES (?,?,?,?,?)',
                                    ('annual_financial_facts', f'{symbol}:{year}:{header}',
                                     'conflicting workbook values',
                                     json.dumps({'kept': current[0], 'rejected': raw,
                                                 'sheet': sheet.title, 'row': source_row},
                                                ensure_ascii=False), imported_at),
                                )
                            continue
                        connection.execute(
                            'INSERT INTO annual_financial_facts VALUES (?,?,?,?,?,?,?,?,?,?,?)',
                            (symbol, year, header, metric_category(header), raw, numeric,
                             source_url, digest, sheet.title, source_row, imported_at),
                        )
                        saved += 1
            quality = dict(status='completed', workbook_rows=workbook_rows,
                           saved_facts=saved, duplicate_facts=duplicates,
                           conflict_facts=conflicts, symbols=len(symbols),
                           min_year=min(years) if years else None,
                           max_year=max(years) if years else None,
                           sheets=len(schemas), schemas=schemas)
            connection.execute(
                'INSERT INTO annual_imports VALUES (?,?,?,?,?,?,?,?,?)',
                (digest, str(workbook), source_url, imported_at, workbook_rows, saved,
                 duplicates, conflicts, json.dumps(quality, ensure_ascii=False)),
            )
        return dict(workbook_sha256=digest, **quality)
    finally:
        connection.close()
